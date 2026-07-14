"""Утилита для формирования xlsx-таблиц с форматированием.

Использование:
    builder = XlsxBuilder("Сотрудники")
    builder.add_sheet(
        headers=["ФИО", "Должность", "Организация"],
        rows=[["Иванов И.И.", "Учитель", "Школа №1"], ...],
    )
    buf = await builder.build()
    # buf — io.BytesIO, готов к отдаче через FastAPI Response
"""

import asyncio
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill(fill_type="solid", fgColor="37474F")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

_CELL_ALIGN = Alignment(vertical="center", wrap_text=False)

_THIN = Side(border_style="thin", color="D0D7DE")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_ALT_FILL = PatternFill(fill_type="solid", fgColor="F5F7F9")


class XlsxBuilder:
    """Строит .xlsx с одним листом: шапка + строки данных."""

    def __init__(self, title: str) -> None:
        self._title = title
        self._headers: list[str] = []
        self._rows: list[list[Any]] = []

    def add_sheet(self, headers: list[str], rows: list[list[Any]]) -> "XlsxBuilder":
        self._headers = headers
        self._rows = rows
        return self

    def _build_sync(self) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = self._title[:31]  # Excel: max 31 символ

        # Заголовки
        for col_idx, header in enumerate(self._headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = _HEADER_ALIGN
            cell.border = _BORDER

        ws.row_dimensions[1].height = 32

        # Данные
        for row_idx, row in enumerate(self._rows, start=2):
            fill = _ALT_FILL if row_idx % 2 == 0 else None
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = _CELL_ALIGN
                cell.border = _BORDER
                if fill:
                    cell.fill = fill
            ws.row_dimensions[row_idx].height = 18

        # Авто-ширина столбцов
        for col_idx, header in enumerate(self._headers, start=1):
            letter = get_column_letter(col_idx)
            # Берём максимум из заголовка и значений в столбце
            max_len = len(str(header))
            for row in self._rows:
                if col_idx - 1 < len(row):
                    max_len = max(max_len, len(str(row[col_idx - 1] or "")))
            # Ограничиваем: минимум 10, максимум 60 символов
            ws.column_dimensions[letter].width = max(10, min(max_len + 4, 60))

        # Закрепить шапку
        ws.freeze_panes = "A2"

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    async def build(self) -> io.BytesIO:
        """Запускает синхронную генерацию в пуле потоков, не блокируя event loop."""
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(None, self._build_sync)
