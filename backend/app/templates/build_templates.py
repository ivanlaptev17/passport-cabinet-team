"""Генерирует исходные .docx-шаблоны с плейсхолдерами docxtpl (Jinja2).

Запуск (один раз, чтобы создать/пересоздать файлы шаблонов):
    docker compose exec backend python -m app.templates.build_templates
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_DIR = Path(__file__).parent

FONT_NAME = "Times New Roman"
BODY_SIZE = Pt(14)
TABLE_SIZE = Pt(12)


def _style_document(doc: Document) -> None:
    """Деловое оформление по умолчанию: Times New Roman, поля под подшивку."""
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    heading = doc.styles["Heading 1"]
    heading.font.name = FONT_NAME
    heading.font.size = Pt(16)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)

    for section in doc.sections:
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)


def _set_cell_font(cell, size=TABLE_SIZE, bold=False) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = FONT_NAME
            run.font.size = size
            run.font.bold = bold


def build_position_certificate():
    doc = Document()
    _style_document(doc)

    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.add_run("{{ logo }}")

    title = doc.add_heading("СПРАВКА", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("Выдана ")
    p.add_run("{{ fio }}").bold = True
    p.add_run(" в том, что он(а) занимает должность ")
    p.add_run("{{ position }}").bold = True
    p.add_run(" в организации ")
    p.add_run("{{ organization_name }}").bold = True
    p.add_run(".")

    doc.add_paragraph()
    note = doc.add_paragraph("Справка выдана для предъявления по месту требования.")
    note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph()
    doc.add_paragraph("Дата выдачи: {{ issue_date }}")

    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph("Директор ____________________ {{ director_name }}")

    doc.save(OUT_DIR / "position_certificate.docx")


def build_responsible_order():
    doc = Document()
    _style_document(doc)

    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.add_run("{{ logo }}")

    org_p = doc.add_paragraph("{{ organization_name }}")
    org_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    org_p.runs[0].bold = True

    title = doc.add_heading("ПРИКАЗ № {{ order_number }}", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date_p = doc.add_paragraph("от {{ order_date }}")
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    body = doc.add_paragraph(
        "В целях обеспечения безопасности образовательного процесса "
        "назначить ответственными лицами:"
    )
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph()

    # 4 строки: заголовок, маркер "{%tr for%}", повторяемая строка данных,
    # маркер "{%tr endfor%}". docxtpl схлопывает строки-маркеры в чистые
    # jinja-теги {% for %} / {% endfor %}, а строка данных между ними
    # повторяется для каждого элемента списка.
    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(1.2), Cm(5.5), Cm(4.5), Cm(5.3)]
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width

    header = table.rows[0].cells
    header[0].text = "№"
    header[1].text = "ФИО"
    header[2].text = "Должность"
    header[3].text = "Зона ответственности"
    for cell in header:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cell_font(cell, bold=True)

    table.rows[1].cells[0].text = "{%tr for p in responsible_persons %}"

    data_row = table.rows[2].cells
    data_row[0].text = "{{ loop.index }}"
    data_row[1].text = "{{ p.fio }}"
    data_row[2].text = "{{ p.position }}"
    data_row[3].text = "{{ p.area }}"
    for cell in data_row:
        _set_cell_font(cell)

    table.rows[3].cells[0].text = "{%tr endfor %}"

    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph("Директор ____________________ {{ director_name }}")

    doc.save(OUT_DIR / "responsible_order.docx")


if __name__ == "__main__":
    build_position_certificate()
    build_responsible_order()
    print("OK: position_certificate.docx, responsible_order.docx")
